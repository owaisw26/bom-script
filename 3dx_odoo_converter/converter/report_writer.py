"""Write the audit workbook used to diagnose warnings and rejected source rows."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import RejectedRow, ValidationWarning


def _format(sheet, table_name: str) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 50
        )
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def write_report(path: Path, *, source_filename: str, source_rows: int, assemblies: int,
                 products: int, component_lines: int, missing_item_numbers: int,
                 assumed_quantities: int, non_released: int, warnings: list[ValidationWarning],
                 rejected_rows: list[RejectedRow], source_headers: list[str],
                 product_mapping: dict[str, str], bom_mapping: dict[str, str]) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    for metric, value in [
        ("Source filename", source_filename), ("Number of source rows", source_rows),
        ("Number of assemblies", assemblies), ("Number of unique products", products),
        ("Number of component lines", component_lines),
        ("Number of missing Enterprise Item Numbers", missing_item_numbers),
        ("Number of assumed quantities", assumed_quantities),
        ("Number of non-released items", non_released),
        ("Number of rejected rows", len(rejected_rows)),
        ("Conversion timestamp", datetime.now(timezone.utc).isoformat()),
    ]:
        summary.append([metric, value])
    _format(summary, "SummaryTable")

    warning_sheet = workbook.create_sheet("Warnings")
    warning_sheet.append(["Source Row", "Warning"])
    for warning in warnings:
        warning_sheet.append([warning.row_number or "", warning.message])
    _format(warning_sheet, "WarningsTable")

    rejected_sheet = workbook.create_sheet("Rejected Rows")
    rejected_sheet.append(["Source Row", "Rejection Reason", *source_headers])
    for rejected in rejected_rows:
        rejected_sheet.append([rejected.row_number or "", rejected.reason,
                               *[rejected.source_row.get(header, "") for header in source_headers]])
    _format(rejected_sheet, "RejectedRowsTable")

    mapping_sheet = workbook.create_sheet("Source Mapping")
    mapping_sheet.append(["Source column", "Odoo field", "Transformation rule"])
    for source, field, rule in [
        ("Enterprise Item Number / Name", product_mapping["external_id"], "Stable 3dx_product_<identifier>"),
        ("Title", product_mapping["name"], "Trimmed source title"),
        ("Enterprise Item Number / Name", product_mapping["internal_reference"], "Engineering number or 3DX-Name"),
        ("Hierarchy parent", bom_mapping["product_external_id"], "Parent product External ID"),
        ("Hierarchy child", bom_mapping["component_external_id"], "Component product External ID"),
        ("Quantity column / default", bom_mapping["component_quantity"], "Configured quantity or assumed default"),
        ("Revision", bom_mapping["external_id"], "Stable 3dx_bom_<parent>_<revision>"),
    ]:
        mapping_sheet.append([source, field, rule])
    _format(mapping_sheet, "SourceMappingTable")
    workbook.save(path)
