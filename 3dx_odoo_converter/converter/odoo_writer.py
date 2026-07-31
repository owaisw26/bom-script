"""Write minimal, Odoo-ready import workbooks."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo


PRODUCT_KEYS = ("external_id", "name", "internal_reference", "product_type", "can_be_sold",
                "can_be_purchased", "uom", "purchase_uom")
BOM_KEYS = ("external_id", "product_external_id", "product_quantity", "bom_type",
            "component_external_id", "component_quantity", "component_uom", "reference")


def load_mapping(path: Path) -> dict[str, dict[str, str]]:
    """Load and validate the configurable Odoo column labels."""
    mapping = json.loads(path.read_text(encoding="utf-8"))
    for section, keys in (("product", PRODUCT_KEYS), ("bom", BOM_KEYS)):
        values = mapping.get(section, {})
        missing = [key for key in keys if not isinstance(values.get(key), str) or not values[key].strip()]
        if missing:
            raise ValueError(f"Field mapping is missing {section} key(s): {', '.join(missing)}")
        if len(set(values[key] for key in keys)) != len(keys):
            raise ValueError(f"Field mapping has duplicate {section} headers")
    return mapping


def _format_sheet(sheet, table_name: str) -> None:
    """Apply simple spreadsheet formatting without adding import-breaking rows."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 50)
        sheet.column_dimensions[column[0].column_letter].width = width
    # Excel tables require at least a header plus one data row.
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def write_import_workbook(path: Path, worksheet_name: str, headers: list[str],
                          rows: list[list[object]], table_name: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = worksheet_name
    sheet.append(headers)
    for row in rows:
        if len(row) != len(headers) or any(value is None or value == "" for value in row):
            raise ValueError("Output workbook contains a blank required cell")
        sheet.append(row)
    _format_sheet(sheet, table_name)
    workbook.save(path)
