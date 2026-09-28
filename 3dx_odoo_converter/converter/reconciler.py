"""Match converted products to product stubs that already exist in Odoo."""

from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .models import OdooBom, OdooProduct, ValidationWarning


@dataclass(frozen=True)
class ExistingProduct:
    external_id: str
    name: str
    internal_reference: str
    row_number: int


@dataclass
class ReconciliationResult:
    matched_products: int
    matched_assemblies: int
    warnings: list[ValidationWarning]
    errors: list[str]


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _key(value: str) -> str:
    return " ".join(value.split()).casefold()


def _header(headers: list[str], *candidates: str) -> str | None:
    """Return the first available display or Odoo technical field name."""
    return next((candidate for candidate in candidates if candidate in headers), None)


def _csv_rows(path: Path) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        headers = reader.fieldnames or []
        return headers, [(row_number, row) for row_number, row in enumerate(reader, start=2)]


def _xlsx_rows(path: Path) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    first_row = next(values, ())
    headers = [_text(value) for value in first_row]
    rows = []
    for row_number, values_row in enumerate(values, start=2):
        rows.append((row_number, dict(zip(headers, values_row))))
    return headers, rows


def load_existing_products(path: Path, product_mapping: dict[str, str]) -> list[ExistingProduct]:
    """Read an Odoo product export used to identify existing product templates."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        headers, rows = _csv_rows(path)
    elif suffix == ".xlsx":
        headers, rows = _xlsx_rows(path)
    else:
        raise ValueError("Existing-products export must be a .csv or .xlsx file")

    external_header = _header(headers, product_mapping["external_id"], "External ID", "id")
    name_header = _header(headers, product_mapping["name"], "Name", "name")
    reference_header = _header(
        headers, product_mapping["internal_reference"], "Internal Reference", "default_code"
    )
    if external_header is None:
        raise ValueError("Existing-products export is missing External ID (id)")
    if name_header is None and reference_header is None:
        raise ValueError(
            "Existing-products export must contain Name (name) or Internal Reference (default_code)"
        )

    products: list[ExistingProduct] = []
    seen_external_ids: set[str] = set()
    for row_number, row in rows:
        external_id = _text(row.get(external_header))
        name = _text(row.get(name_header)) if name_header else ""
        internal_reference = _text(row.get(reference_header)) if reference_header else ""
        if not external_id and not name and not internal_reference:
            continue
        if not external_id:
            raise ValueError(f"Existing-products export row {row_number} has no External ID")
        if not name and not internal_reference:
            raise ValueError(f"Existing-products export row {row_number} has no matching value")
        if external_id in seen_external_ids:
            raise ValueError(f"Existing-products export has duplicate External ID '{external_id}'")
        seen_external_ids.add(external_id)
        products.append(ExistingProduct(external_id, name, internal_reference, row_number))
    if not products:
        raise ValueError("Existing-products export contains no product rows")
    return products


def _index(existing: list[ExistingProduct], attribute: str) -> dict[str, list[ExistingProduct]]:
    values: dict[str, list[ExistingProduct]] = defaultdict(list)
    for product in existing:
        value = _key(getattr(product, attribute))
        if value:
            values[value].append(product)
    return values


def reconcile_existing_products(products: list[OdooProduct], boms: list[OdooBom],
                                existing: list[ExistingProduct]) -> ReconciliationResult:
    """Reuse Odoo External IDs and reject any assembly that would be duplicated."""
    by_reference = _index(existing, "internal_reference")
    by_name = _index(existing, "name")
    assembly_external_ids = {bom.product_external_id for bom in boms}
    replacements: dict[str, str] = {}
    claimed_existing_ids: dict[str, str] = {}
    warnings: list[ValidationWarning] = []
    errors: list[str] = []
    matched_assemblies = 0

    for product in products:
        reference_matches = by_reference.get(_key(product.internal_reference), [])
        name_matches = by_name.get(_key(product.name), [])
        chosen: ExistingProduct | None = None
        matched_by = ""

        if len(reference_matches) > 1:
            errors.append(
                f"Existing Odoo products have an ambiguous Internal Reference '{product.internal_reference}'"
            )
        elif len(reference_matches) == 1:
            chosen = reference_matches[0]
            matched_by = "Internal Reference"

        if len(name_matches) > 1 and chosen is None:
            errors.append(f"Existing Odoo products have an ambiguous Name '{product.name}'")
        elif len(name_matches) == 1:
            name_match = name_matches[0]
            if chosen is not None and chosen.external_id != name_match.external_id:
                errors.append(
                    f"Product '{product.name}' matches different Odoo stubs by Name and Internal Reference"
                )
            elif chosen is None:
                chosen = name_match
                matched_by = "Name"

        is_assembly = product.external_id in assembly_external_ids
        if chosen is None:
            if is_assembly and not reference_matches and not name_matches:
                errors.append(
                    f"Assembly '{product.name}' was not found in the existing-products export; "
                    "refusing to create a duplicate"
                )
            continue

        other_product = claimed_existing_ids.get(chosen.external_id)
        if other_product and other_product != product.name:
            errors.append(
                f"Converted products '{other_product}' and '{product.name}' both match "
                f"Odoo External ID '{chosen.external_id}'"
            )
            continue
        claimed_existing_ids[chosen.external_id] = product.name
        replacements[product.external_id] = chosen.external_id
        if is_assembly:
            matched_assemblies += 1
        warnings.append(ValidationWarning(
            None,
            f"Matched existing Odoo product '{product.name}' by {matched_by}; "
            f"reused External ID '{chosen.external_id}'",
        ))

    if errors:
        return ReconciliationResult(0, 0, warnings, errors)

    for product in products:
        product.external_id = replacements.get(product.external_id, product.external_id)
    for bom in boms:
        bom.product_external_id = replacements.get(bom.product_external_id, bom.product_external_id)
    return ReconciliationResult(len(replacements), matched_assemblies, warnings, [])
