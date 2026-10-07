import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
HEADER = "Level,Title,Enterprise Item Number,Revision,Maturity State,Name,Type\n"
MANUAL_HEADER = "Project ID,Parent Reference,Internal Reference,Name,Quantity,Unit of Measure,Status\n"


def run_converter(source: Path, output: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run([sys.executable, "convert_3dx_to_odoo.py", str(source),
                           "--output-dir", str(output), *args], cwd=PROJECT,
                          text=True, capture_output=True, check=False)


def test_manual_parts_join_only_released_installations(tmp_path: Path) -> None:
    source = tmp_path / "installations.csv"
    source.write_text(HEADER +
        "0,Ready Assembly,ASM-READY,A.1,Released,ready,Physical Product\n"
        "1,Made Part,PART-1,A.1,Released,part,Physical Product\n"
        "0,Immature Assembly,ASM-DRAFT,A.1,In Work,draft,Physical Product\n"
        "1,Future Part,PART-2,A.1,Released,future,Physical Product\n", encoding="utf-8")
    manual = tmp_path / "procurement.csv"
    manual.write_text(MANUAL_HEADER +
        "P-123,ASM-READY,BUY-1,Fastener,3,Units,Ready\n"
        "P-123,ASM-READY,BUY-1,Fastener,2,Units,Ready\n"
        "P-123,ASM-DRAFT,BUY-2,Future Bracket,1,Units,Ready\n"
        "P-123,ASM-READY,BUY-3,Unapproved,1,Units,Draft\n", encoding="utf-8")
    output = tmp_path / "out"
    completed = run_converter(source, output, "--project-id", "P-123",
                              "--procurement-csv", str(manual))
    assert completed.returncode == 0, completed.stderr
    products = list(load_workbook(output / "odoo_products_import.xlsx").active.values)
    boms = list(load_workbook(output / "odoo_bom_import.xlsx").active.values)
    report = load_workbook(output / "conversion_report.xlsx")
    assert {row[2] for row in products[1:]} == {"ASM-READY", "PART-1", "BUY-1"}
    assert {row[2]: row[5] for row in products[1:]} == {
        "ASM-READY": "False", "PART-1": "False", "BUY-1": "True"}
    assert {row[5]: row[6] for row in boms[1:]} == {"PART-1": 1, "BUY-1": 5}
    assert all(row[8].startswith("P-123 | ") for row in boms[1:])
    assert [row[8] for row in list(report["Procurement"].values)[1:]] == [
        "Exported", "Exported", "Held", "Held"]


def test_drawing_export_fails_closed_and_removes_stale_imports(tmp_path: Path) -> None:
    source = tmp_path / "drawing.csv"
    source.write_text(HEADER +
        "0,Drawing,ASM,A.1,Released,root,Drawing\n", encoding="utf-8")
    output = tmp_path / "out"
    output.mkdir()
    (output / "odoo_products_import.xlsx").write_text("stale")
    completed = run_converter(source, output)
    assert completed.returncode == 1
    assert "export installations instead" in completed.stderr
    assert not (output / "odoo_products_import.xlsx").exists()


def test_manual_part_requires_valid_quantity(tmp_path: Path) -> None:
    source = tmp_path / "installations.csv"
    source.write_text(HEADER +
        "0,Assembly,ASM,A.1,Released,root,Physical Product\n", encoding="utf-8")
    manual = tmp_path / "procurement.csv"
    manual.write_text(MANUAL_HEADER +
        "P-1,ASM,BUY-1,Fastener,0,Units,Ready\n", encoding="utf-8")
    output = tmp_path / "out"
    completed = run_converter(source, output, "--project-id", "P-1",
                              "--procurement-csv", str(manual))
    assert completed.returncode == 1
    assert "positive finite number" in completed.stderr
    assert not (output / "odoo_products_import.xlsx").exists()
