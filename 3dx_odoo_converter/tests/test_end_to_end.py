import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
SAMPLE = PROJECT.parent / "ASM-DOORS INSTALLATION.csv"


def test_supplied_csv_converts_to_expected_workbooks(tmp_path: Path) -> None:
    completed = subprocess.run(
        [sys.executable, "convert_3dx_to_odoo.py", str(SAMPLE), "--output-dir", str(tmp_path)],
        cwd=PROJECT, text=True, capture_output=True, check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert "Product rows exported: 5" in completed.stdout
    products = load_workbook(tmp_path / "odoo_products_import.xlsx").active
    boms = load_workbook(tmp_path / "odoo_bom_import.xlsx").active
    report = load_workbook(tmp_path / "conversion_report.xlsx")
    assert products.title == "Products"
    assert boms.title == "Bills of Materials"
    assert products.max_row == 6
    assert boms.max_row == 5
    assert products[1][0].value == "External ID"
    assert [cell.value for cell in boms[1]] == [
        "External ID", "Product/External ID", "Quantity", "BoM Type",
        "BoM Lines/External ID", "BoM Lines/Component", "BoM Lines/Quantity",
        "BoM Lines/Unit", "Reference",
    ]
    product_references = {row[2] for row in products.iter_rows(min_row=2, values_only=True)}
    component_references = {row[5] for row in boms.iter_rows(min_row=2, values_only=True)}
    assert component_references <= product_references
    line_external_ids = [row[4] for row in boms.iter_rows(min_row=2, values_only=True)]
    assert len(line_external_ids) == len(set(line_external_ids))
    assert all(value.startswith("3dx_bom_") for value in line_external_ids)
    assert list(products.tables) == ["ProductsImport"]
    assert list(boms.tables) == ["BillsOfMaterialsImport"]
    assert report.sheetnames == ["Summary", "Warnings", "Rejected Rows", "Source Mapping"]


def test_existing_assembly_stub_is_updated_instead_of_duplicated(tmp_path: Path) -> None:
    existing = tmp_path / "existing_products.csv"
    existing.write_text(
        "External ID,Name,Internal Reference\n"
        "__export__.product_template_assembly,ASM-DOORS INSTALLATION,\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    completed = subprocess.run(
        [
            sys.executable, "convert_3dx_to_odoo.py", str(SAMPLE),
            "--existing-products", str(existing), "--output-dir", str(output),
        ],
        cwd=PROJECT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "Existing assemblies matched: 1" in completed.stdout
    products = load_workbook(output / "odoo_products_import.xlsx").active
    boms = load_workbook(output / "odoo_bom_import.xlsx").active
    product_rows = list(products.iter_rows(min_row=2, values_only=True))
    bom_rows = list(boms.iter_rows(min_row=2, values_only=True))
    assembly = next(row for row in product_rows if row[1] == "ASM-DOORS INSTALLATION")
    assert assembly[0] == "__export__.product_template_assembly"
    assert {row[1] for row in bom_rows} == {"__export__.product_template_assembly"}


def test_reconciliation_refuses_to_create_a_missing_assembly(tmp_path: Path) -> None:
    existing = tmp_path / "existing_products.csv"
    existing.write_text(
        "External ID,Name,Internal Reference\n"
        "__export__.product_template_other,OTHER ASSEMBLY,OTHER-1\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    completed = subprocess.run(
        [
            sys.executable, "convert_3dx_to_odoo.py", str(SAMPLE),
            "--existing-products", str(existing), "--output-dir", str(output),
        ],
        cwd=PROJECT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 1
    assert "refusing to create a duplicate" in completed.stderr
    assert not (output / "odoo_products_import.xlsx").exists()
    assert not (output / "odoo_bom_import.xlsx").exists()
    assert (output / "conversion_report.xlsx").is_file()
